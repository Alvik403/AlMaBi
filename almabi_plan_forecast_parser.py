from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from almabi_excel_utils import MONTH_NAMES, normalize_header, normalize_text, parse_amount
from almabi_pipeline import Fact

PLAN_FORECAST_HEADERS = {
    "direction": "направление",
    "project_group": "группа проектов",
    "project": "проект",
    "tax_type": "вид налогообложения",
    "section": "раздел",
    "cost_kind_future": "вид себестоимости (будущее)",
    "cost_kind": "вид себестоимости",
    "amount": "сумма",
    "month": "месяц",
    "scenario": "план/прогноз",
}

KPI_NAMES = frozenset(
    {
        "Выручка",
        "Себестоимость",
        "Коммерческие расходы",
        "Управленческие расходы",
        "Прочие доходы",
        "Прочие расходы",
        "Налоги",
    }
)

EXPENSE_KPIS = frozenset(
    {
        "Себестоимость",
        "Коммерческие расходы",
        "Управленческие расходы",
        "Прочие расходы",
        "Налоги",
    }
)

KNOWN_COST_SECTIONS = frozenset(
    {
        "Материальные затраты",
        "ФОТ",
        "Аренда (прямые)",
        "Амортизация",
        "Прочие производственные расходы",
        "Общепроизводственные затраты",
    }
)


@dataclass(frozen=True)
class PlanForecastParseResult:
    plan_facts: list[Fact]
    forecast_facts: list[Fact]
    warnings: list[str]


def _normalize_dimension(value: object) -> str:
    text = normalize_text(value)
    if not text or text.isdigit():
        return ""
    return text


def _normalize_tax_type(value: object) -> str:
    text = normalize_text(value)
    if not text:
        return "Общие условия налогообложения"
    if "льгот" in text.casefold():
        return "Льготное налогообложение"
    return text


def _normalize_month(value: object) -> str | None:
    text = normalize_text(value)
    if not text:
        return None
    if text in MONTH_NAMES.values():
        return text
    lowered = text.casefold().rstrip(".")
    for full in MONTH_NAMES.values():
        if full.casefold() == lowered:
            return full
        if full.casefold().startswith(lowered) or lowered.startswith(full[:3].casefold()):
            return full
    return None


def _normalize_scenario(value: object) -> str | None:
    text = normalize_text(value).casefold()
    if text.startswith("план"):
        return "План"
    if text.startswith("прогн"):
        return "Прогноз"
    return None


def _resolve_kpi(section: str, cost_kind: str, cost_kind_future: str) -> str | None:
    for candidate in (cost_kind, cost_kind_future):
        if not candidate:
            continue
        if candidate in KPI_NAMES:
            return candidate
        lowered = candidate.casefold()
        if "выруч" in lowered:
            return "Выручка"
        if "себест" in lowered:
            return "Себестоимость"
        if "коммерч" in lowered:
            return "Коммерческие расходы"
        if "управлен" in lowered:
            return "Управленческие расходы"
        if "проч" in lowered and "доход" in lowered:
            return "Прочие доходы"
        if "проч" in lowered and "расход" in lowered:
            return "Прочие расходы"
        if "налог" in lowered:
            return "Налоги"

    section_lower = section.casefold()
    if "доход" in section_lower:
        return "Прочие доходы" if cost_kind and "проч" in cost_kind.casefold() else "Выручка"
    if "расход" in section_lower:
        if cost_kind and "коммерч" in cost_kind.casefold():
            return "Коммерческие расходы"
        if cost_kind and "управлен" in cost_kind.casefold():
            return "Управленческие расходы"
        if cost_kind and "проч" in cost_kind.casefold():
            return "Прочие расходы"
        return "Себестоимость"
    return None


def _resolve_cost_section(kpi_l1: str, cost_kind: str, cost_kind_future: str) -> str:
    if kpi_l1 != "Себестоимость":
        return ""
    for candidate in (cost_kind, cost_kind_future):
        if candidate in KNOWN_COST_SECTIONS:
            return candidate
        if candidate in KPI_NAMES:
            continue
    return "Общепроизводственные затраты"


def _signed_amount(kpi_l1: str, amount: float) -> float:
    if kpi_l1 in EXPENSE_KPIS:
        return -abs(amount)
    return abs(amount)


def _find_header_row(sheet) -> tuple[int, dict[str, int]]:
    for row_index in range(1, min(sheet.max_row, 20) + 1):
        mapping: dict[str, int] = {}
        for col_index in range(1, sheet.max_column + 1):
            header = normalize_header(sheet.cell(row_index, col_index).value)
            for key, expected in PLAN_FORECAST_HEADERS.items():
                if header == expected:
                    mapping[key] = col_index
        if {"amount", "month", "scenario"}.issubset(mapping):
            return row_index, mapping
    raise ValueError("Не найдена строка заголовков формы план/прогноз (нужны колонки: Сумма, Месяц, План/прогноз).")


def _cell(row: tuple[object, ...], mapping: dict[str, int], key: str) -> object:
    index = mapping.get(key)
    if index is None:
        return None
    return row[index - 1]


def parse_plan_forecast_workbook(path: Path) -> PlanForecastParseResult:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        header_row, mapping = _find_header_row(sheet)
        plan_facts: list[Fact] = []
        forecast_facts: list[Fact] = []
        warnings: list[str] = []

        for row_index, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, values_only=True),
            start=header_row + 1,
        ):
            amount = parse_amount(_cell(row, mapping, "amount"))
            if not amount:
                continue

            month = _normalize_month(_cell(row, mapping, "month"))
            scenario = _normalize_scenario(_cell(row, mapping, "scenario"))
            section = normalize_text(_cell(row, mapping, "section"))
            cost_kind = normalize_text(_cell(row, mapping, "cost_kind"))
            cost_kind_future = normalize_text(_cell(row, mapping, "cost_kind_future"))
            kpi_l1 = _resolve_kpi(section, cost_kind, cost_kind_future)

            if not month:
                warnings.append(f"Строка {row_index}: пропущена — не указан месяц.")
                continue
            if not scenario:
                warnings.append(f"Строка {row_index}: пропущена — не указан План/прогноз.")
                continue
            if not kpi_l1:
                warnings.append(
                    f"Строка {row_index}: не удалось определить KPI "
                    f"(раздел={section!r}, вид={cost_kind!r})."
                )
                continue

            signed = _signed_amount(kpi_l1, amount)
            fact = Fact(
                kpi_l1=kpi_l1,
                month=month,
                amount_buh=signed,
                amount_nu=signed,
                direction=_normalize_dimension(_cell(row, mapping, "direction")),
                project_group=_normalize_dimension(_cell(row, mapping, "project_group")),
                project=_normalize_dimension(_cell(row, mapping, "project")),
                cost_section=_resolve_cost_section(kpi_l1, cost_kind, cost_kind_future),
                expense_article=cost_kind if kpi_l1 in {"Прочие доходы", "Прочие расходы"} else "",
                tax_type=_normalize_tax_type(_cell(row, mapping, "tax_type")),
            )
            if scenario == "План":
                plan_facts.append(fact)
            else:
                forecast_facts.append(fact)

        return PlanForecastParseResult(plan_facts=plan_facts, forecast_facts=forecast_facts, warnings=warnings)
    finally:
        workbook.close()


def validate_plan_forecast_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        _find_header_row(workbook.active)
    finally:
        workbook.close()
