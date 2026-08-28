"""Generate AlMaBi network infrastructure checklist (Word) for IT/SB approval."""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
REF_TEMPLATE = Path(r"C:\Users\Flirp\Downloads\Чек_лист_сетевой_инфраструктуры_ОПиОП_BI.docx")
OUT = Path(r"C:\Users\Flirp\Downloads\Чек_лист_сетевой_инфраструктуры_AlMaBi.docx")
DIAGRAMS = ROOT / "docs" / "network_checklist" / "diagrams"
IMG_NETWORK = ROOT / "docs" / "network_checklist" / "network_33.png"
IMG_DATAFLOW = ROOT / "docs" / "network_checklist" / "dataflow_35.png"

P = {
    "desc": 1,
    "launch": 2,
    "contact": 3,
    "dcod": 6,
    "cloud": 7,
    "office": 8,
    "workplace": 9,
    "contractor": 10,
    "host_existing": 12,
    "host_new": 13,
    "containers": 14,
    "vlan_yes": 16,
    "vlan_no": 17,
    "dns_yes": 19,
    "dns_no": 20,
    "all_internal": 23,
    "dept": 24,
    "users": 25,
    "internal_systems": 26,
    "public_internet": 28,
    "vpn_users": 29,
    "partners": 30,
    "external_api": 31,
    "img_network": 33,
    "src_data": 39,
    "dst_data": 40,
    "transform": 41,
    "paths": 42,
    "admin": 43,
    "auth_yes": 45,
    "auth_no": 46,
    "outbound_yes": 54,
    "outbound_no": 55,
    "ssl_yes": 58,
    "ssl_existing": 59,
    "ssl_no": 60,
    "tls_yes": 62,
    "tls_no": 63,
    "pdn_yes": 65,
    "pdn_no": 66,
    "extra": 69,
    "append_network": 72,
    "append_dataflow": 73,
    "img_dataflow": 74,
}


def set_text(doc: Document, index: int, text: str) -> None:
    doc.paragraphs[index].text = text


def set_paragraph_image(paragraph, image_path: Path, width_inches: float = 6.2) -> None:
    if not image_path.is_file():
        return
    element = paragraph._element
    for child in list(element):
        if child.tag.endswith("r"):
            element.remove(child)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def export_diagrams() -> None:
    exports = [
        (DIAGRAMS / "network_33.mmd", IMG_NETWORK),
        (DIAGRAMS / "dataflow_35.mmd", IMG_DATAFLOW),
    ]
    npx = "npx.cmd" if sys.platform == "win32" else "npx"
    for src, dst in exports:
        if not src.is_file():
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        cmd = [
            npx,
            "-y",
            "@mermaid-js/mermaid-cli",
            "-i",
            str(src),
            "-o",
            str(dst),
            "-b",
            "transparent",
            "-w",
            "1400",
        ]
        subprocess.run(cmd, check=True, cwd=ROOT, shell=sys.platform == "win32")


def fill_paragraphs(doc: Document) -> None:
    set_text(
        doc,
        P["desc"],
        "1.1. Описание проекта и его целей: внутренний веб-дашборд Al Ma BI (AlMaBi) "
        "для построения управленческого P&L по Excel-выгрузкам 1С (бухрегистр, "
        "реализация, себестоимость; опционально план/прогноз)",
    )
    set_text(doc, P["launch"], "1.2. Планируемая дата запуска проекта: 30.09.2026")
    set_text(
        doc,
        P["contact"],
        "1.3. Контактное лицо по проекту (ФИО, номер телефона, e-mail): "
        "Попов Алексей Викторович, +79934058605, Alepopov@alabuga.ru",
    )
    set_text(doc, P["dcod"], "Локальный ЦОД Компании — целевой вариант; площадку согласует ИТ")
    set_text(doc, P["cloud"], "Облако (указать какое) — не требуется приложением")
    set_text(doc, P["office"], "Серверное помещение в офисе — не планируется")
    set_text(
        doc,
        P["workplace"],
        "Рабочее место пользователя — серверные компоненты не размещаются; используется только браузер",
    )
    set_text(doc, P["contractor"], "У подрядчиков/партнеров — не размещается")
    set_text(
        doc,
        P["host_existing"],
        "Используем существующие (указать hostname): возможно; hostname и IP production-хоста "
        "предоставляет ИТ",
    )
    set_text(
        doc,
        P["host_new"],
        "Требуются новые сервера/BM (указать кол-во, версия ОС): только если нет подходящего хоста. "
        "Нужен Linux-хост/ВМ с Docker Engine и Docker Compose v2. Лимиты compose: web 1,5 CPU/1 ГБ, "
        "nginx 0,5 CPU/256 МБ; ресурсы хоста определяет ИТ",
    )
    set_text(
        doc,
        P["containers"],
        "Требуются контейнеры (Docker, K8s): Docker Compose prod (docker-compose.prod.yml): "
        "nginx (80→8080, 443→8443) и web (8000 только во внутренней Docker-сети)",
    )
    set_text(doc, P["vlan_yes"], "Да — только по решению ИТ/СБ; приложение отдельную VLAN не требует")
    set_text(
        doc,
        P["vlan_no"],
        "Нет (Будет использоваться имеющийся): серверный сегмент с доступом из согласованных LAN/VPN",
    )
    set_text(
        doc,
        P["dns_yes"],
        "Да (указать желаемое имя): требуется корпоративное DNS-имя; FQDN задаётся в "
        "ALMABI_HOST/ALLOWED_HOSTS и SAN TLS-сертификата",
    )
    set_text(doc, P["dns_no"], "Нет, используем существующий — только если ИТ выделит подходящее DNS-имя")
    set_text(doc, P["all_internal"], "Все пользователи — нет; доступ только согласованным сотрудникам с личной УЗ")
    set_text(
        doc,
        P["dept"],
        "Конкретные отделы (перечислить): перечень подразделений предоставляет владелец системы",
    )
    set_text(
        doc,
        P["users"],
        "Конкретные пользователи или группы (указать): список сотрудников и роли "
        "viewer/uploader/admin предоставляет владелец",
    )
    set_text(
        doc,
        P["internal_systems"],
        "Другие внутренние системы/серверы (указать): мониторинг опрашивает /health и /ready через nginx",
    )
    set_text(doc, P["public_internet"], "Все пользователи Интернета (публичный сервис) — нет")
    set_text(
        doc,
        P["vpn_users"],
        "Только сотрудники компании (через VPN) — да, если нужен доступ вне корпоративной LAN",
    )
    set_text(doc, P["partners"], "Только определенные партнеры/контрагенты (предоставить список IP-адресов): —")
    set_text(doc, P["external_api"], "Внешние API/сервисы (перечислить): —")

    set_paragraph_image(doc.paragraphs[P["img_network"]], IMG_NETWORK)

    set_text(
        doc,
        P["src_data"],
        "Источник данных (системы и приложения, которые генерируют данные) "
        "(например: базы данных): Excel .xlsx — выгрузки 1С; загрузка uploader/admin",
    )
    set_text(
        doc,
        P["dst_data"],
        "Назначение данных (системы и приложения которые потребляют данные) "
        "например: CRM-системы): веб-дашборд AlMaBi — P&L, drill-down, PQ-таблицы, KPI-отчёты",
    )
    set_text(
        doc,
        P["transform"],
        "Трансформация данных (процессы, которые изменяют формат или структуру данных, "
        "чтобы сделать их совместимыми с пунктом назначения или более полезными для анализа "
        "(например: очистка, агрегация): проверка XLSX, парсинг openpyxl, PQ pipeline, "
        "агрегации P&L в памяти; SQLite только для УЗ и сессий",
    )
    set_text(
        doc,
        P["paths"],
        "Пути потоков данных (маршруты, по которым данные перемещаются между компонентами): "
        "браузер → HTTPS:443 → nginx → HTTP web:8000 → uploads/almabi/users/{user_id}/ → "
        "обработка в памяти → HTML/JSON; runtime/auth.sqlite3; logs",
    )
    set_text(
        doc,
        P["admin"],
        "3.6. Кто будет администрировать систему? Попов Алексей Викторович — контакт по приложению; "
        "ИТ назначает администратора хоста, Docker, TLS и backup",
    )
    set_text(
        doc,
        P["auth_yes"],
        "Да — локальные УЗ (SQLite auth.sqlite3), RBAC viewer/uploader/admin, Argon2id, CSRF; "
        "управление через scripts/manage_users.py и /admin/users",
    )
    set_text(doc, P["auth_no"], "Нет — не применяется для production")
    set_text(doc, P["outbound_yes"], "Да (описать куда именно): —")
    set_text(
        doc,
        P["outbound_no"],
        "Нет — в runtime нет внешних API/CDN; только same-origin fetch в браузере",
    )
    set_text(
        doc,
        P["ssl_yes"],
        "Да (указать тип: public/self-signed/internal CA): тип CA и FQDN согласуются; "
        "nginx использует ops/tls/fullchain.pem и privkey.pem",
    )
    set_text(doc, P["ssl_existing"], "Нет, используется существующий — если ИТ выдаст подходящий сертификат")
    set_text(doc, P["ssl_no"], "Нет — не допускается для production")
    set_text(
        doc,
        P["tls_yes"],
        "Да — TLS 1.2/1.3 пользователь ↔ nginx (№1); nginx → web по HTTP во internal Docker-сети (№3)",
    )
    set_text(doc, P["tls_no"], "Нет — не применяется для внешнего пользовательского трафика")
    set_text(
        doc,
        P["pdn_yes"],
        "Да — возможны финансовые данные и ПДн в выгрузках 1С; решение принимают владелец и СБ",
    )
    set_text(doc, P["pdn_no"], "Нет — только если владелец и СБ подтвердят отсутствие ПДн")
    set_text(
        doc,
        P["extra"],
        " порт web:8000 не публиковать на хост; AUTH_ENABLED=true; SESSION_HTTPS_ONLY=true; "
        "секреты в .env.prod; backup uploads/runtime/logs; лимит файла 50 МиБ, квота 500 МiБ; "
        "retention uploads 90 дней, logs 30 дней",
    )
    set_text(doc, P["append_network"], "Схема сетевого взаимодействия (п. 3.3)")
    set_text(doc, P["append_dataflow"], "Схема потоков данных (п. 3.5)")
    set_paragraph_image(doc.paragraphs[P["img_dataflow"]], IMG_DATAFLOW)


def fill_tables(doc: Document) -> None:
    network_rows = [
        ("Наименование", "№ связи", "SRC", "DST", "Protocol", "Port", "Примечание"),
        (
            "Пользователь (браузер, LAN/VPN)",
            "1",
            "IP АРМ / VPN-пул — уточнить",
            "IP production-хоста — уточнить",
            "TCP",
            "443",
            "HTTPS, HSTS, rate limit, allow-list частных сетей",
        ),
        (
            "Пользователь (браузер, LAN/VPN)",
            "2",
            "IP АРМ / VPN-пул — уточнить",
            "IP production-хоста — уточнить",
            "TCP",
            "80",
            "HTTP → HTTPS redirect",
        ),
        (
            "nginx (контейнер)",
            "3",
            "Docker-сеть backend",
            "web:8000",
            "TCP",
            "8000",
            "Внутренняя Docker-сеть; порт не на хосте",
        ),
        (
            "Мониторинг",
            "4",
            "IP мониторинга — уточнить",
            "IP production-хоста — уточнить",
            "TCP",
            "443",
            "GET /health, /ready через HTTPS",
        ),
        ("", "", "", "", "", "", ""),
    ]
    access_rows = [
        ("Роль", "Ресурс 1", "Ресурс 2 ", "Ресурс 3"),
        ("Роль", "Дашборд и отчёты", "Загрузка Excel", "УЗ и Swagger"),
        ("admin", "чтение", "запись", "управление УЗ, Swagger"),
        ("uploader", "чтение", "запись", "нет"),
        ("viewer", "чтение", "нет", "нет"),
    ]
    for table_index, rows in enumerate((network_rows, access_rows)):
        table = doc.tables[table_index]
        while len(table.rows) > 0:
            table._tbl.remove(table.rows[-1]._tr)
        for row_data in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row_data):
                cell.text = value


def main() -> None:
    if not REF_TEMPLATE.is_file():
        raise FileNotFoundError(f"Reference template not found: {REF_TEMPLATE}")
    try:
        export_diagrams()
    except subprocess.CalledProcessError as exc:
        print(f"Warning: diagram export failed ({exc}); continuing without PNG refresh")
    doc = Document(str(REF_TEMPLATE))
    fill_paragraphs(doc)
    fill_tables(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
