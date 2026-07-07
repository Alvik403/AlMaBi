from __future__ import annotations

from pathlib import Path

from almabi_pipeline_audit import PipelineAuditLog
from almabi_test_pipeline import build_test_facts, run_test_pipeline
from tests.test_almabi_exports import (
    create_buh_workbook,
    create_cost_workbook,
    create_realization_workbook,
)
from almabi_export_parsers import parse_exports


def test_audit_logs_revenue_and_cost_lines(tmp_path: Path):
    paths = {
        "buh": tmp_path / "buh.xlsx",
        "realization": tmp_path / "realization.xlsx",
        "cost": tmp_path / "cost.xlsx",
    }
    create_buh_workbook(paths["buh"])
    create_realization_workbook(paths["realization"])
    create_cost_workbook(paths["cost"])

    audit = PipelineAuditLog()
    build_test_facts(parse_exports(paths), audit=audit)

    assert len(audit.revenue_lines) == 1
    assert len(audit.cost_lines) == 1
    assert audit.revenue_lines[0]["analytics"]["direction"] == "Услуги"
    assert audit.revenue_lines[0]["realization_join"] is not None
    assert audit.cost_lines[0]["cost_join"] is not None


def test_audit_writes_report_files(tmp_path: Path):
    paths = {
        "buh": tmp_path / "buh.xlsx",
        "realization": tmp_path / "realization.xlsx",
        "cost": tmp_path / "cost.xlsx",
    }
    create_buh_workbook(paths["buh"])
    create_realization_workbook(paths["realization"])
    create_cost_workbook(paths["cost"])

    logs_dir = tmp_path / "logs"
    pipeline = run_test_pipeline(paths, logs_dir=logs_dir, write_audit=True)

    assert pipeline.audit_path is not None
    assert pipeline.audit_path.exists()
    assert (logs_dir / "almabi_test" / "audit-latest.json").exists()
    assert pipeline.audit.to_dict()["summary"]["revenue_lines"] == 1
    assert pipeline.audit.to_dict()["summary"]["cost_lines"] == 1


def test_audit_logs_other_sections(tmp_path: Path):
    from openpyxl import load_workbook

    paths = {
        "buh": tmp_path / "buh.xlsx",
        "realization": tmp_path / "realization.xlsx",
        "cost": tmp_path / "cost.xlsx",
    }
    create_buh_workbook(paths["buh"])
    create_realization_workbook(paths["realization"])
    create_cost_workbook(paths["cost"])

    workbook = load_workbook(paths["buh"])
    sheet = workbook.active
    row_idx = sheet.max_row
    sheet.insert_rows(row_idx)
    values = [
        "Операция 0001 от 15.03.2026",
        "76.09",
        "",
        "",
        "91.01",
        "",
        "Проценты полученные",
        "15.03.2026",
        "",
        400_000,
        0,
        400_000,
    ]
    for col, value in enumerate(values, start=1):
        sheet.cell(row_idx, col, value)
    workbook.save(paths["buh"])

    audit = PipelineAuditLog()
    build_test_facts(parse_exports(paths), audit=audit)

    assert len(audit.section_lines.get("Коммерческие расходы", [])) == 1
    assert len(audit.section_lines.get("Прочие доходы", [])) == 1
    payload = audit.to_dict()
    assert payload["summary"]["other_section_lines"]["Коммерческие расходы"] == 1
    assert payload["summary"]["other_section_lines"]["Прочие доходы"] == 1


def test_audit_detects_duplicate_buh_rows(tmp_path: Path):
    from openpyxl import Workbook

    from tests.test_almabi_exports import _pad_rows

    buh_path = tmp_path / "buh.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    headers = [
        "Документ",
        "Счет Дт",
        "Счет Кт",
        "Субконто1 Кт",
        "Дата",
        "Сумма",
        "Сумма НУ Дт",
        "Сумма НУ Кт",
    ]
    sheet.append(headers)
    duplicate_row = [
        "Реализация товаров и услуг 00АМ-000017 от 31.01.2026 21:00:00",
        "62.01",
        "90.01.3",
        "",
        "15.01.2026",
        1_000_000,
        0,
        1_000_000,
    ]
    sheet.append(duplicate_row)
    sheet.append(duplicate_row)
    sheet.append(["Итого"])
    workbook.save(buh_path)

    realization_path = tmp_path / "realization.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    create_realization_workbook(realization_path)
    create_cost_workbook(cost_path)

    audit = PipelineAuditLog()
    build_test_facts(parse_exports({"buh": buh_path, "realization": realization_path, "cost": cost_path}), audit=audit)

    assert len(audit.duplicate_buh_rows) == 1
    assert audit.duplicate_buh_rows[0]["count"] == 2
    assert "Выручка" in audit.duplicate_buh_rows[0]["sections"]
    assert len(audit.duplicate_facts) == 1
    assert audit.duplicate_facts[0]["count"] == 2
    assert len(audit.duplicate_facts[0]["occurrences"]) == 2
    assert all("bi_path" in item for item in audit.duplicate_facts[0]["occurrences"])
