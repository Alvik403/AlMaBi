from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from almabi_pq_common import resolve_header_index
from almabi_pq_cost import (
    COST_AMOUNT_HEADERS,
    _resolve_cost_amount_index,
    build_pq_cost_table,
    table_summary,
)
from almabi_pq_projects import PQ_PROJECTS_SECTION, build_pq_projects_table
from tests.test_almabi_exports import create_cost_workbook, create_realization_workbook


def _pad_rows(sheet, count: int) -> None:
    for _ in range(count):
        sheet.append([])


def create_wide_cost_workbook(path: Path) -> None:
    """Широкая выгрузка как в 1С (скрин пользователя)."""
    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 5)
    headers = [
        "Продукция",
        "Счет",
        "Статья калькуляции",
        "Документ отгрузки",
        "Количество продаж",
        "Количество затрат",
        "Себестоимость (бухг. учет)",
        "% от себестоимости",
        "Себестоимость на единицу",
        "Прямые расходы",
        "Стоимость единицы затрат",
        "Выручка (бухг. учет)",
        "Валовая прибыль (убыток)",
        "ОХР (26)",
        "Расходы на продажу (44)",
        "Прочие расходы (91)",
        "Себестоимость полная",
    ]
    sheet.append(headers)
    sheet.append(
        [
            "Лицензия ПО",
            20,
            "Сырье и материалы",
            "Реализация 00АМ-000300 от 31.01.2026",
            1,
            1,
            400_000,
            99.9,
            400_000,
            0,
            0,
            1_000_000,
            600_000,
            0,
            0,
            0,
            999_999_999,
        ]
    )
    _pad_rows(sheet, 38)
    workbook.save(path)


def test_resolve_header_index_prefers_buhg_not_percent():
    headers = [
        "Продукция",
        "Себестоимость (бухг. учет)",
        "% от себестоимости",
        "Себестоимость полная",
    ]
    index = resolve_header_index(headers, COST_AMOUNT_HEADERS)
    assert index == 1
    assert headers[index] == "Себестоимость (бухг. учет)"


def test_resolve_cost_amount_index_prefers_regl_over_buhg():
    headers = [
        "Продукция",
        "Себестоимость (регл. учет)",
        "Себестоимость (бухг. учет)",
        "% от себестоимости",
    ]
    index = _resolve_cost_amount_index(headers)
    assert index == 1
    assert headers[index] == "Себестоимость (регл. учет)"


def test_build_pq_cost_table_prefers_regl_amount(tmp_path: Path):
    path = tmp_path / "regl-buhg-cost.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 5)
    sheet.append(
        [
            "Продукция",
            "Счет",
            "Статья калькуляции",
            "Документ отгрузки",
            "Количество продаж",
            "Себестоимость (регл. учет)",
            "Себестоимость (бухг. учет)",
        ]
    )
    sheet.append(["Лицензия ПО", 20, "Сырье и материалы", "Реализация 001", 1, 300_000, 400_000])
    _pad_rows(sheet, 38)
    workbook.save(path)

    rows = build_pq_cost_table(path)

    assert len(rows) == 1
    assert rows[0]["Сумма"] == 300_000


def test_join_projects_uses_first_match_when_duplicates(tmp_path: Path):
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    create_wide_cost_workbook(cost_path)

    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 7)
    sheet.append(
        [
            "Заказ клиента / Реализация",
            "Номенклатура",
            "Проект",
            "Группа проектов",
            "Направление",
            "Выручка",
        ]
    )
    document = "Реализация 00АМ-000300 от 31.01.2026"
    sheet.append([document, "Лицензия ПО", "Проект A", "Группа A", "Направление A", 1_000_000])
    sheet.append([document, "Лицензия ПО", "Проект B", "Группа B", "Направление B", 500_000])
    _pad_rows(sheet, 1)
    workbook.save(realization_path)

    rows = build_pq_cost_table(cost_path, projects_path=realization_path)
    cost_row = next(row for row in rows if row["Раздел"] == "Материальные затраты")

    assert cost_row["Проект"] == "Проект A"
    assert cost_row["Направление"] == "Направление A"


def test_wide_cost_workbook_uses_buhg_amount_not_full(tmp_path: Path):
    path = tmp_path / "wide-cost.xlsx"
    create_wide_cost_workbook(path)

    rows = build_pq_cost_table(path)

    assert len(rows) == 1
    assert rows[0]["Сумма"] == 400_000
    assert rows[0]["Раздел"] == "Материальные затраты"


def test_build_pq_cost_table_with_projects(tmp_path: Path):
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    create_wide_cost_workbook(cost_path)
    create_realization_workbook(
        realization_path,
        document="Реализация 00АМ-000300 от 31.01.2026",
    )

    rows = build_pq_cost_table(cost_path, projects_path=realization_path)

    assert len(rows) >= 2
    cost_row = next(row for row in rows if row["Раздел"] == "Материальные затраты")
    revenue_row = next(row for row in rows if row["Раздел"] == "Выручка")
    assert cost_row["Сумма"] == 400_000
    assert cost_row["Основной раздел"] == "Расходы"
    assert revenue_row["Сумма"] == 1_000_000
    assert revenue_row["Основной раздел"] == "Доходы"
    assert revenue_row["Номенклатура"] == "Лицензия ПО"


def test_build_pq_cost_table_compact_workbook(tmp_path: Path):
    path = tmp_path / "cost.xlsx"
    create_cost_workbook(path)

    rows = build_pq_cost_table(path)

    assert len(rows) == 1
    assert rows[0]["Сумма"] == 400_000
    assert rows[0]["Раздел"] == "Материальные затраты"


def test_build_pq_cost_table_header_on_row_4(tmp_path: Path):
    path = tmp_path / "cost-row4.xlsx"
    create_cost_workbook(path, header_pad=3)

    rows = build_pq_cost_table(path)

    assert len(rows) == 1
    assert rows[0]["Сумма"] == 400_000
    assert rows[0]["Раздел"] == "Материальные затраты"


def test_build_pq_cost_table_header_on_row_6(tmp_path: Path):
    path = tmp_path / "cost-row6.xlsx"
    create_cost_workbook(path, header_pad=5)

    rows = build_pq_cost_table(path)

    assert len(rows) == 1
    assert rows[0]["Сумма"] == 400_000


def test_build_pq_projects_table(tmp_path: Path):
    path = tmp_path / "realization.xlsx"
    create_realization_workbook(path)

    rows = build_pq_projects_table(path)

    assert len(rows) == 1
    assert rows[0]["Раздел"] == PQ_PROJECTS_SECTION
    assert rows[0]["Номенклатура"] == "Лицензия ПО"
    assert rows[0]["Сумма"] == 1_000_000


def test_cost_table_summary():
    rows = [{"Сумма": 100.0}, {"Сумма": 50.0}]
    summary = table_summary(rows)
    assert summary["row_count"] == 2
    assert summary["total_amount"] == 150.0
