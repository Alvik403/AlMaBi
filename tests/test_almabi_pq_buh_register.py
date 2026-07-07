from __future__ import annotations

from pathlib import Path

from almabi_pq_buh_register import (
    build_pq_buh_register_table,
    classify_pq_buh_section,
    classify_pq_main_section,
    table_summary,
)
from tests.test_almabi_exports import create_buh_workbook, create_cost_workbook, create_realization_workbook


def test_classify_pq_buh_section_names_match_power_query():
    assert classify_pq_buh_section("90.07.1", "44") == "Расходы на продажу"
    assert classify_pq_buh_section("90.08.1", "26") == "Управленческие расходы"
    assert classify_pq_main_section("Расходы на продажу") == "Расходы"


def test_build_pq_buh_register_with_cost_and_revenue_joins(tmp_path: Path):
    buh_path = tmp_path / "buh.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    create_buh_workbook(buh_path)
    create_cost_workbook(cost_path)
    create_realization_workbook(realization_path)

    rows = build_pq_buh_register_table(
        buh_path,
        cost_path=cost_path,
        revenue_path=realization_path,
        projects_path=realization_path,
    )

    assert len(rows) >= 3
    cost_row = next(row for row in rows if row["Раздел"] == "Себестоимость")
    revenue_row = next(row for row in rows if row["Раздел"] == "Выручка")
    selling_row = next(row for row in rows if row["Раздел"] == "Расходы на продажу")

    assert cost_row["Основной раздел"] == "Расходы"
    assert cost_row["Сумма БУ"] == -400_000
    assert cost_row["Себестоимость.Раздел"] == "Материальные затраты"
    assert cost_row["Себестоимость.Сумма"] == 400_000

    assert revenue_row["Основной раздел"] == "Доходы"
    assert revenue_row["Сумма БУ"] == 1_000_000
    assert revenue_row["Проект"] == "Обслуживание Долго"
    assert revenue_row["Направление"] == "Услуги"

    assert selling_row["Сумма БУ"] == -50_000
    assert selling_row["Статья дохода расхода"] == "Реклама"


def test_buh_table_summary_uses_amount_buh():
    rows = [{"Сумма БУ": 100.0}, {"Сумма БУ": -40.0}]
    summary = table_summary(rows)
    assert summary["total_amount"] == 60.0
    assert summary["amount_column"] == "Сумма БУ"


def test_buh_keeps_rows_without_amounts_and_document(tmp_path: Path):
    from openpyxl import Workbook

    from tests.test_almabi_exports import _pad_rows

    path = tmp_path / "buh-sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    headers = ["Документ", "Счет Дт", "Счет Кт", "Сумма", "Сумма НУ Дт", "Сумма НУ Кт", "Дата"]
    sheet.append(headers)
    sheet.append(["", "62.01", "51", 0, 0, 0, "15.01.2026"])
    sheet.append(["Док без сумм", "62.01", "51", None, None, None, "16.01.2026"])
    sheet.append(["Итого"])
    workbook.save(path)

    rows = build_pq_buh_register_table(path)

    assert len(rows) == 2
    assert rows[0]["Документ"] == ""
    assert rows[1]["Документ"] == "Док без сумм"


def test_buh_cost_join_duplicates_rows_like_power_query(tmp_path: Path):
    from openpyxl import Workbook

    from tests.test_almabi_exports import _pad_rows

    buh_path = tmp_path / "buh.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    document = "Реализация товаров и услуг 00АМ-000017 от 31.01.2026 21:00:00"

    create_realization_workbook(realization_path, document=document)

    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    sheet.append(["Документ", "Счет Дт", "Счет Кт", "Субконто1 Кт", "Сумма", "Сумма НУ Дт", "Сумма НУ Кт", "Дата"])
    sheet.append([document, "90.02.1", "43", "Лицензия ПО", 400_000, 400_000, 0, "15.01.2026"])
    sheet.append(["Итого"])
    workbook.save(buh_path)

    cost_wb = Workbook()
    cost_sheet = cost_wb.active
    _pad_rows(cost_sheet, 5)
    cost_sheet.append(
        [
            "Продукция",
            "Счет",
            "Статья калькуляции",
            "Документ отгрузки",
            "Количество продаж",
            "Себестоимость (бухг. учет)",
        ]
    )
    cost_sheet.append(["Лицензия ПО", "20", "Сырье и материалы", document, 1, 200_000])
    cost_sheet.append(["Лицензия ПО", "20", "Оплата труда", document, 1, 200_000])
    _pad_rows(cost_sheet, 38)
    cost_wb.save(cost_path)

    rows = build_pq_buh_register_table(
        buh_path,
        cost_path=cost_path,
        revenue_path=realization_path,
        projects_path=realization_path,
    )

    cost_lines = [row for row in rows if row["Раздел"] == "Себестоимость"]
    assert len(cost_lines) == 2
    assert {row["Себестоимость.Сумма"] for row in cost_lines} == {200_000}
    assert sum(row["Сумма БУ"] for row in cost_lines) == -400_000
