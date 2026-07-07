from __future__ import annotations

from pathlib import Path

from almabi_test_pipeline import (
    _coalesce_pq_value,
    build_test_facts,
    classify_main_section,
)
from tests.test_almabi_exports import (
    create_buh_workbook,
    create_cost_workbook,
    create_realization_workbook,
)
from almabi_export_parsers import parse_exports


def test_classify_main_section():
    assert classify_main_section("Выручка") == "Доходы"
    assert classify_main_section("Прочие доходы") == "Доходы"
    assert classify_main_section("Себестоимость") == "Расходы"
    assert classify_main_section("Коммерческие расходы") == "Расходы"


def test_coalesce_pq_prefers_single_source():
    assert _coalesce_pq_value("Услуги", "") == "Услуги"
    assert _coalesce_pq_value("", "Услуги") == "Услуги"
    assert _coalesce_pq_value("Услуги", "Производство") == "Производство"
    assert _coalesce_pq_value("", "") == ""


def test_revenue_direction_from_realization_join(tmp_path: Path):
    paths = {
        "buh": tmp_path / "buh.xlsx",
        "realization": tmp_path / "realization.xlsx",
        "cost": tmp_path / "cost.xlsx",
    }
    create_buh_workbook(paths["buh"])
    create_realization_workbook(paths["realization"])
    create_cost_workbook(paths["cost"])

    result = build_test_facts(parse_exports(paths))
    revenue = [fact for fact in result.facts if fact.kpi_l1 == "Выручка"]
    assert revenue
    assert revenue[0].direction == "Услуги"
    assert revenue[0].project_group == "Обслуживание"
    assert revenue[0].project == "Обслуживание Долго"


def test_cost_direction_when_document_text_differs(tmp_path: Path):
    """Бухрегистр и себестоимость с разным текстом документа, но одним номером 00АМ-000017."""
    buh_doc = "Проводка по реализации 00АМ-000017 от 31.01.2026"
    cost_doc = "Реализация товаров и услуг 00АМ-000017 от 31.01.2026 21:00:00"

    buh_path = tmp_path / "buh.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    create_buh_workbook(buh_path, document=buh_doc)
    create_cost_workbook(cost_path, document=cost_doc)
    create_realization_workbook(realization_path, document=cost_doc)

    result = build_test_facts(
        parse_exports({"buh": buh_path, "realization": realization_path, "cost": cost_path}),
        cost_path=cost_path,
        projects_path=realization_path,
    )
    cost = [fact for fact in result.facts if fact.kpi_l1 == "Себестоимость"]
    assert cost
    assert cost[0].direction == "Услуги"
    assert cost[0].project_group == "Обслуживание"
    assert cost[0].project == "Обслуживание Долго"


def test_cost_uses_cost_file_amount(tmp_path: Path):
    paths = {
        "buh": tmp_path / "buh.xlsx",
        "realization": tmp_path / "realization.xlsx",
        "cost": tmp_path / "cost.xlsx",
    }
    create_buh_workbook(paths["buh"])
    create_realization_workbook(paths["realization"])
    create_cost_workbook(paths["cost"])

    result = build_test_facts(parse_exports(paths))
    cost = [fact for fact in result.facts if fact.kpi_l1 == "Себестоимость"]
    assert len(cost) == 1
    assert cost[0].amount_buh == -400_000
    assert cost[0].direction == "Услуги"
    assert cost[0].project_group == "Обслуживание"
    assert cost[0].project == "Обслуживание Долго"


def test_other_income_has_no_direction_without_join(tmp_path: Path):
    from openpyxl import Workbook

    from tests.test_almabi_exports import _pad_rows

    buh_path = tmp_path / "buh.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    headers = [
        "Документ",
        "Счет Дт",
        "Счет Кт",
        "Субконто1 Кт",
        "Дата",
        "Сумма",
        "Сумма НУ Дт",
        "Сумма НУ Кт",
    ]
    sheet.append(headers)
    sheet.append(
        [
            "Операция 0001 от 15.03.2026",
            "76.09",
            "91.01",
            "Проценты полученные",
            "15.03.2026",
            400_000,
            0,
            400_000,
        ]
    )
    sheet.append(["Итого"])
    workbook.save(buh_path)

    realization_path = tmp_path / "realization.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    create_realization_workbook(realization_path, document="Другой документ")
    create_cost_workbook(cost_path, document="Другой документ")

    result = build_test_facts(parse_exports({"buh": buh_path, "realization": realization_path, "cost": cost_path}))
    other = [fact for fact in result.facts if fact.kpi_l1 == "Прочие доходы"]
    assert len(other) == 1
    assert other[0].direction == "Без направления"
    assert other[0].amount_buh == 400_000


def test_cost_join_expands_like_power_query(tmp_path: Path):
    from openpyxl import Workbook

    from almabi_pq_buh_register import build_pq_buh_register_table
    from tests.test_almabi_exports import _pad_rows

    buh_path = tmp_path / "buh.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    document = "Реализация товаров и услуг 00АМ-000017 от 31.01.2026 21:00:00"

    create_realization_workbook(realization_path, document=document)

    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    sheet.append(["Документ", "Счет Дт", "Счет Кт", "Субконто1 Кт", "Сумма", "Сумма НУ Дт", "Сумма НУ Кт", "Дата"])
    sheet.append([document, "90.02.1", "43", "Лицензия ПО", 400_000, 400_000, 0, "15.01.2026"])
    sheet.append(["Итого"])
    workbook.save(buh_path)

    cost_wb = Workbook()
    cost_sheet = cost_wb.active
    _pad_rows(cost_sheet, 5)
    cost_sheet.append(
        [
            "Продукция",
            "Счет",
            "Статья калькуляции",
            "Документ отгрузки",
            "Количество продаж",
            "Себестоимость (бухг. учет)",
        ]
    )
    cost_sheet.append(["Лицензия ПО", "20", "Сырье и материалы", document, 1, 200_000])
    cost_sheet.append(["Лицензия ПО", "20", "Оплата труда", document, 1, 200_000])
    _pad_rows(cost_sheet, 38)
    cost_wb.save(cost_path)

    pq_rows = build_pq_buh_register_table(
        buh_path,
        cost_path=cost_path,
        revenue_path=realization_path,
        projects_path=realization_path,
    )
    pq_cost_total = sum(row["Сумма БУ"] for row in pq_rows if row["Раздел"] == "Себестоимость")

    result = build_test_facts(parse_exports({"buh": buh_path, "realization": realization_path, "cost": cost_path}))
    test_cost_total = sum(fact.amount_buh for fact in result.facts if fact.kpi_l1 == "Себестоимость")

    assert len([fact for fact in result.facts if fact.kpi_l1 == "Себестоимость"]) == 2
    assert test_cost_total == pq_cost_total == -400_000


def test_cost_without_join_matches_pq_sign(tmp_path: Path):
    """Строки бухрегистра без join к себестоимости: «Сумма БУ» положительная, как в PQ."""
    from openpyxl import Workbook

    from almabi_pq_buh_register import build_pq_buh_register_table
    from tests.test_almabi_exports import _pad_rows, create_realization_workbook

    buh_path = tmp_path / "buh.xlsx"
    cost_path = tmp_path / "cost.xlsx"
    realization_path = tmp_path / "realization.xlsx"
    document = "Реализация товаров и услуг 00АМ-000099 от 31.01.2026 21:00:00"

    create_realization_workbook(realization_path, document=document)

    workbook = Workbook()
    sheet = workbook.active
    _pad_rows(sheet, 8)
    sheet.append(["Документ", "Счет Дт", "Счет Кт", "Субконто1 Кт", "Сумма", "Сумма НУ Дт", "Сумма НУ Кт", "Дата"])
    sheet.append([document, "90.02.1", "43", "Нет в файле себестоимости", 500_000, 500_000, 0, "15.01.2026"])
    sheet.append(["Итого"])
    workbook.save(buh_path)

    cost_wb = Workbook()
    cost_sheet = cost_wb.active
    _pad_rows(cost_sheet, 5)
    cost_sheet.append(["Продукция", "Счет", "Статья калькуляции", "Документ отгрузки", "Количество продаж", "Себестоимость (бухг. учет)"])
    cost_sheet.append(["Другая номенклатура", "20", "Сырье и материалы", document, 1, 100_000])
    _pad_rows(cost_sheet, 38)
    cost_wb.save(cost_path)

    exports = parse_exports({"buh": buh_path, "realization": realization_path, "cost": cost_path})
    pq_rows = build_pq_buh_register_table(
        buh_path,
        cost_path=cost_path,
        revenue_path=realization_path,
        projects_path=realization_path,
    )
    pq_total = sum(float(row["Сумма БУ"]) for row in pq_rows if row["Раздел"] == "Себестоимость")
    test_total = sum(
        fact.amount_buh
        for fact in build_test_facts(exports).facts
        if fact.kpi_l1 == "Себестоимость"
    )

    assert pq_total == test_total == 500_000
